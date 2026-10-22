"""Workbook inspection and compatibility preflight.

`inspect_workbook` reads structure, formulas, cached values, names, tables,
hidden content and workbook features, then classifies what the prototype can
safely do with the file. Function-name scanning is a preflight heuristic only;
the round-trip check in `engine.validate_round_trip` is what actually enables
editing.
"""

from __future__ import annotations

import datetime as dt
import io
import multiprocessing
import re
from typing import Any

from openpyxl import load_workbook
from openpyxl.formula.tokenizer import Token, Tokenizer
from openpyxl.worksheet.formula import ArrayFormula, DataTableFormula

from app.errors import AppError
from app.workbook.package import package_features

# Formula families verified by compatibility fixtures. Extend only together with new fixtures.
SUPPORTED_FUNCTIONS = frozenset({"SUM", "IF"})
VOLATILE_FUNCTIONS = frozenset({"NOW", "TODAY", "RAND", "RANDBETWEEN", "RANDARRAY", "OFFSET", "INDIRECT",
                                "CELL", "INFO"})
DYNAMIC_ARRAY_FUNCTIONS = frozenset({"FILTER", "SORT", "SORTBY", "UNIQUE", "SEQUENCE", "XLOOKUP", "XMATCH",
                                     "LET", "LAMBDA", "MAP", "REDUCE", "SCAN", "MAKEARRAY", "BYROW", "BYCOL",
                                     "VSTACK", "HSTACK", "TOCOL", "TOROW", "WRAPROWS", "WRAPCOLS", "TAKE",
                                     "DROP", "CHOOSECOLS", "CHOOSEROWS", "EXPAND", "TEXTSPLIT"})
DATE_FUNCTIONS = frozenset({"DATE", "DATEDIF", "DATEVALUE", "DAY", "DAYS", "EDATE", "EOMONTH", "MONTH", "YEAR",
                            "NETWORKDAYS", "WORKDAY", "WEEKDAY", "WEEKNUM", "YEARFRAC", "TIME", "HOUR", "MINUTE"})
_EXTERNAL_REF = re.compile(r"\[[^\]]+\]")
_SIMPLE_RANGE = re.compile(r"^(?:'[^']+'|[A-Za-z0-9_.]+)!\$?[A-Z]{1,3}\$?\d+(?::\$?[A-Z]{1,3}\$?\d+)?$")


class InspectionFailed(AppError):
    def __init__(self, code: str, message: str):
        super().__init__(code, message, status_code=422)


def _json_value(value: Any) -> Any:
    if value is None or isinstance(value, (bool, int, float, str)):
        return value
    if isinstance(value, (dt.datetime, dt.date, dt.time)):
        return {"type": "date", "iso": value.isoformat()}
    return str(value)


def _functions_in(formula: str) -> tuple[set[str], bool]:
    """Return function names used in a formula and whether it references an external workbook."""
    functions: set[str] = set()
    external = False
    try:
        tokens = Tokenizer(formula).items
    except Exception:  # unparseable formula text
        return {"<UNPARSEABLE>"}, False
    for token in tokens:
        if token.type == Token.FUNC and token.subtype == Token.OPEN:
            name = token.value[:-1].upper()
            for prefix in ("_XLFN._XLWS.", "_XLFN.", "_XLWS."):
                if name.startswith(prefix):
                    name = name[len(prefix):]
            functions.add(name)
        elif token.type == Token.OPERAND and token.subtype == Token.RANGE and _EXTERNAL_REF.search(token.value):
            external = True
    return functions, external


def inspect_bytes(data: bytes, max_cells: int) -> dict[str, Any]:
    """Inspect a validated .xlsx package. Runs in a child process (see `inspect_workbook`)."""
    features = package_features(data)
    wb = load_workbook(io.BytesIO(data), data_only=False, keep_links=True)
    wv = load_workbook(io.BytesIO(data), data_only=True, keep_links=True)

    sheets = []
    formulas: dict[str, dict[str, str]] = {}
    cached: dict[str, dict[str, Any]] = {}
    values: dict[str, dict[str, Any]] = {}
    functions: set[str] = set()
    array_formulas = data_tables = external_formula_refs = 0
    total_cells = 0

    formats: dict[str, dict[str, str]] = {}
    for ws in wb.worksheets:
        wsv = wv[ws.title]
        sheet_formulas: dict[str, str] = {}
        sheet_cached: dict[str, Any] = {}
        sheet_values: dict[str, Any] = {}
        sheet_formats: dict[str, str] = {}
        formats[ws.title] = sheet_formats
        populated = 0
        for row in ws.iter_rows():
            for cell in row:
                value = cell.value
                if value is None:
                    continue
                populated += 1
                if cell.number_format and cell.number_format != "General":
                    sheet_formats[cell.coordinate] = cell.number_format
                total_cells += 1
                if total_cells > max_cells:
                    raise InspectionFailed("too_many_cells",
                                           f"The workbook has more than {max_cells:,} populated cells.")
                if isinstance(value, ArrayFormula):
                    array_formulas += 1
                    sheet_formulas[cell.coordinate] = value.text or ""
                elif isinstance(value, DataTableFormula):
                    data_tables += 1
                    sheet_formulas[cell.coordinate] = "{=TABLE()}"
                elif cell.data_type == "f" and isinstance(value, str):
                    sheet_formulas[cell.coordinate] = value
                    used, external = _functions_in(value)
                    functions |= used
                    external_formula_refs += int(external)
                else:
                    sheet_values[cell.coordinate] = _json_value(value)
                    continue
                sheet_cached[cell.coordinate] = _json_value(wsv[cell.coordinate].value)
        formulas[ws.title] = sheet_formulas
        cached[ws.title] = sheet_cached
        values[ws.title] = sheet_values
        hidden_rows = [idx for idx, dim in ws.row_dimensions.items() if dim.hidden]
        hidden_cols = [key for key, dim in ws.column_dimensions.items() if dim.hidden]
        sheets.append({
            "name": ws.title,
            "state": ws.sheet_state,
            "dimensions": ws.dimensions,
            "populated_cells": populated,
            "formula_cells": len(sheet_formulas),
            "hidden_rows": hidden_rows[:200],
            "hidden_columns": hidden_cols[:200],
            "protected": bool(ws.protection.sheet),
            "merged_ranges": len(ws.merged_cells.ranges),
            "data_validations": len(ws.data_validations.dataValidation),
            "conditional_formats": len(list(ws.conditional_formatting)),
            "tables": [{"name": t.name, "ref": t.ref} for t in ws.tables.values()],
            "has_comments": any(c.comment for r in ws.iter_rows() for c in r if c.comment),
        })
    for ws in wb.chartsheets:
        sheets.append({"name": ws.title, "state": ws.sheet_state, "chartsheet": True, "populated_cells": 0,
                       "formula_cells": 0, "tables": []})

    defined_names = []
    complex_names = 0
    for name, dn in wb.defined_names.items():
        text = dn.attr_text or ""
        simple = bool(_SIMPLE_RANGE.match(text))
        complex_names += int(not simple)
        defined_names.append({"name": name, "value": text, "scope": "workbook", "simple_range": simple})
    for ws in wb.worksheets:
        for name, dn in ws.defined_names.items():
            text = dn.attr_text or ""
            simple = bool(_SIMPLE_RANGE.match(text))
            complex_names += int(not simple)
            defined_names.append({"name": name, "value": text, "scope": ws.title, "simple_range": simple})

    calc = wb.calculation
    inspection = {
        "sheets": sheets,
        "defined_names": defined_names,
        "formulas": formulas,
        "cached_values": cached,
        "input_values": values,
        "number_formats": formats,
        "functions_used": sorted(functions),
        "populated_cells": total_cells,
        "features": {
            **features,
            "drawings": features["drawings"] and not _only_vml_drawings(data),
            "array_formulas": array_formulas,
            "data_tables": data_tables,
            "external_formula_references": external_formula_refs,
            "complex_named_formulas": complex_names,
            "iterative_calculation": bool(calc is not None and calc.iterate),
            "chartsheets": len(wb.chartsheets),
        },
    }
    inspection["compatibility"] = classify(inspection)
    return inspection


def _only_vml_drawings(data: bytes) -> bool:
    import zipfile
    names = [n.lower() for n in zipfile.ZipFile(io.BytesIO(data)).namelist() if n.lower().startswith("xl/drawings/")]
    return all(n.endswith(".vml") or "/_rels/" in n and n.endswith(".vml.rels") for n in names)


def classify(inspection: dict[str, Any]) -> dict[str, Any]:
    """Preflight classification. Blocking issues prevent editing; inspection stays possible."""
    f = inspection["features"]
    blocking: list[dict[str, str]] = []
    warnings: list[dict[str, str]] = []

    def block(code: str, message: str) -> None:
        blocking.append({"code": code, "message": message})

    def warn(code: str, message: str) -> None:
        warnings.append({"code": code, "message": message})

    if f["external_links"] or f["external_formula_references"]:
        block("external_references", "The workbook refers to other files; results could depend on data we cannot see.")
    if f["data_connections"] or f["query_tables"] or f["power_query_model"]:
        block("data_connections", "The workbook has data connections or Power Query content, which are not refreshed.")
    if f["pivot_tables"] or f["slicers"]:
        block("pivot_tables", "Pivot tables and slicers are not supported and could be damaged by editing.")
    if f["charts"] or f["drawings"] or f["chartsheets"]:
        block("charts_or_drawings", "Charts, images, or drawings would not be preserved by the editing pipeline.")
    if f["embedded_objects"]:
        block("embedded_objects", "Embedded objects or controls are not supported.")
    if f["array_formulas"]:
        block("array_formulas", "Array formulas are not yet covered by compatibility tests.")
    if f["data_tables"]:
        block("data_tables", "What-if data tables are not supported.")
    if f["iterative_calculation"]:
        block("iterative_calculation", "Iterative (circular) calculation is not supported.")
    if f["complex_named_formulas"]:
        block("complex_named_formulas", "Some named ranges contain formulas rather than plain cell references.")

    used = set(inspection["functions_used"])
    dynamic = sorted(used & DYNAMIC_ARRAY_FUNCTIONS)
    volatile = sorted(used & VOLATILE_FUNCTIONS)
    dates = sorted(used & DATE_FUNCTIONS)
    cube = sorted(n for n in used if n.startswith("CUBE"))
    unknown = sorted(used - SUPPORTED_FUNCTIONS - set(dynamic) - set(volatile) - set(dates) - set(cube))
    if dynamic:
        block("dynamic_arrays", f"Dynamic-array functions are not supported yet: {', '.join(dynamic)}.")
    if volatile:
        block("volatile_functions", f"Volatile functions give results that change on every calculation: "
                                    f"{', '.join(volatile)}.")
    if dates:
        block("date_functions", f"Date-dependent functions are not yet covered by tests: {', '.join(dates)}.")
    if cube:
        block("cube_functions", f"Cube functions are not supported: {', '.join(cube)}.")
    if unknown:
        block("untested_functions", f"These functions are not yet covered by compatibility tests: "
                                    f"{', '.join(unknown)}.")

    for sheet in inspection["sheets"]:
        if sheet.get("state") in ("hidden", "veryHidden"):
            warn("hidden_sheet", f"Sheet '{sheet['name']}' is hidden. It stays hidden and is excluded from AI context.")
        if sheet.get("hidden_rows") or sheet.get("hidden_columns"):
            warn("hidden_rows_columns", f"Sheet '{sheet['name']}' has hidden rows or columns; they stay hidden.")
        if sheet.get("protected"):
            warn("protected_sheet", f"Sheet '{sheet['name']}' is protected. Its cells cannot be mapped as editable "
                                    "inputs; protection is never bypassed.")
        if sheet.get("conditional_formats"):
            warn("conditional_formatting", f"Sheet '{sheet['name']}' has conditional formatting, which may not "
                                           "be preserved exactly.")
        if sheet.get("has_comments"):
            warn("comments", f"Sheet '{sheet['name']}' has comments; their text is treated as untrusted data.")

    total_formulas = sum(len(v) for v in inspection["formulas"].values())
    missing_cache = sum(1 for sheet in inspection["cached_values"].values() for v in sheet.values() if v is None)
    if total_formulas and missing_cache:
        warn("missing_cached_values", f"{missing_cache} of {total_formulas} formula cells have no stored result, "
                                      "so calculation fidelity cannot be verified against Excel.")
    return {
        "preflight": "blocked" if blocking else "candidate",
        "blocking": blocking,
        "warnings": warnings,
        "formula_cells": total_formulas,
        "formula_cells_without_cache": missing_cache,
        "supported_functions": sorted(SUPPORTED_FUNCTIONS),
    }


def _child(data: bytes, max_cells: int, conn) -> None:  # pragma: no cover - runs in a child process
    try:
        conn.send(("ok", inspect_bytes(data, max_cells)))
    except AppError as exc:
        conn.send(("app_error", (exc.code, exc.message)))
    except Exception as exc:  # parser failures on malformed files
        conn.send(("error", type(exc).__name__))
    finally:
        conn.close()


def inspect_workbook(data: bytes, *, max_cells: int, timeout: float) -> dict[str, Any]:
    """Inspect in a separate process so a pathological file cannot hang the server."""
    ctx = multiprocessing.get_context("forkserver")
    parent, child = ctx.Pipe(duplex=False)
    proc = ctx.Process(target=_child, args=(data, max_cells, child), daemon=True)
    proc.start()
    child.close()
    try:
        if not parent.poll(timeout):
            raise InspectionFailed("inspection_timeout", "Inspecting the workbook took too long.")
        status, payload = parent.recv()
    except EOFError as exc:
        raise InspectionFailed("inspection_failed", "The workbook could not be read.") from exc
    finally:
        if proc.is_alive():
            proc.kill()
        proc.join(5)
        parent.close()
    if status == "ok":
        return payload
    if status == "app_error":
        raise InspectionFailed(*payload)
    raise InspectionFailed("inspection_failed", "The workbook could not be read; it may be damaged.")
