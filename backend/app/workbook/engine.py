"""WorkbookEngine interface and the LibreOffice implementation.

Extension point: a native-Excel engine (for workbooks that need Excel-specific
behaviour) would implement the same protocol, e.g. by automating Excel on a
Windows host. It is intentionally not built for the MVP.
"""

from __future__ import annotations

import io
import time
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Protocol

from openpyxl import load_workbook

from app.calc import protocol
from app.errors import AppError
from app.workbook import compare
from app.workbook.inspect import inspect_workbook

FORMULA_PREFIXES = ("=", "+", "-", "@")


@dataclass(frozen=True)
class CellWrite:
    sheet: str
    address: str
    value: Any  # float | int | str | bool


@dataclass
class RecalcOutcome:
    ok: bool
    data: bytes | None
    error_code: str | None = None
    error_message: str | None = None
    diagnostics: dict[str, Any] | None = None


class WriteRejected(AppError):
    def __init__(self, code: str, message: str):
        super().__init__(code, message, status_code=422)


class WorkbookEngine(Protocol):
    def inspect(self, data: bytes) -> dict[str, Any]: ...

    def apply_inputs(self, data: bytes, writes: list[CellWrite]) -> bytes: ...

    def submit_recalculation(self, job_id: str, data: bytes) -> None: ...

    def collect_recalculation(self, job_id: str) -> RecalcOutcome | None: ...

    def recalculate(self, job_id: str, data: bytes, wait_seconds: float) -> RecalcOutcome | None: ...

    def read_results(self, data: bytes, cells: list[tuple[str, str]]) -> dict[tuple[str, str], Any]: ...

    def validate_round_trip(self, original: dict[str, Any], recalculated: dict[str, Any],
                            reference_values: dict[tuple[str, str], Any] | None) -> dict[str, Any]: ...


class LibreOfficeEngine:
    """openpyxl for inspection and controlled input edits; LibreOffice (in the worker) for recalculation."""

    name = "libreoffice"

    def __init__(self, *, jobs_dir: Path, calc_timeout: float, max_cells: int, inspection_timeout: float,
                 tolerance: compare.Tolerance):
        self.jobs_dir = jobs_dir
        self.calc_timeout = calc_timeout
        self.max_cells = max_cells
        self.inspection_timeout = inspection_timeout
        self.tolerance = tolerance

    def inspect(self, data: bytes) -> dict[str, Any]:
        return inspect_workbook(data, max_cells=self.max_cells, timeout=self.inspection_timeout)

    def apply_inputs(self, data: bytes, writes: list[CellWrite]) -> bytes:
        wb = load_workbook(io.BytesIO(data), keep_links=True)
        for write in writes:
            if write.sheet not in wb.sheetnames:
                raise WriteRejected("unknown_sheet", f"Sheet '{write.sheet}' does not exist.")
            ws = wb[write.sheet]
            cell = ws[write.address]
            if cell.data_type == "f":
                raise WriteRejected("formula_cell", f"{write.sheet}!{write.address} contains a formula and "
                                                    "cannot be overwritten.")
            if ws.protection.sheet and cell.protection.locked:
                raise WriteRejected("protected_cell", f"{write.sheet}!{write.address} is protected.")
            if any(write.address in r and write.address != r.start_cell.coordinate for r in ws.merged_cells.ranges):
                raise WriteRejected("merged_cell", f"{write.sheet}!{write.address} is inside a merged range.")
            value = write.value
            if isinstance(value, str) and value.lstrip().startswith(FORMULA_PREFIXES):
                raise WriteRejected("formula_injection", "Text inputs may not start with =, +, - or @.")
            if not isinstance(value, (int, float, str, bool)) or isinstance(value, float) and value != value:
                raise WriteRejected("invalid_value", "Unsupported input value.")
            cell.value = value
        if wb.calculation is not None:
            wb.calculation.fullCalcOnLoad = True
        out = io.BytesIO()
        wb.save(out)
        return out.getvalue()

    def submit_recalculation(self, job_id: str, data: bytes) -> None:
        protocol.submit(self.jobs_dir, protocol.RecalcRequest(job_id=job_id, timeout_seconds=self.calc_timeout), data)

    def collect_recalculation(self, job_id: str) -> RecalcOutcome | None:
        found = protocol.read_result(self.jobs_dir, job_id)
        if found is None:
            return None
        result, data = found
        if result.ok and data is None:
            return RecalcOutcome(False, None, "missing_output", "The calculation produced no workbook.")
        return RecalcOutcome(result.ok, data, result.error_code, result.error_message, result.diagnostics)

    def recalculate(self, job_id: str, data: bytes, wait_seconds: float) -> RecalcOutcome | None:
        self.submit_recalculation(job_id, data)
        deadline = time.monotonic() + wait_seconds
        while time.monotonic() < deadline:
            outcome = self.collect_recalculation(job_id)
            if outcome is not None:
                return outcome
            time.sleep(0.2)
        return None

    def read_results(self, data: bytes, cells: list[tuple[str, str]]) -> dict[tuple[str, str], Any]:
        wb = load_workbook(io.BytesIO(data), data_only=True)
        return {(sheet, addr): wb[sheet][addr].value if sheet in wb.sheetnames else None for sheet, addr in cells}

    def validate_round_trip(self, original: dict[str, Any], recalculated: dict[str, Any],
                            reference_values: dict[tuple[str, str], Any] | None) -> dict[str, Any]:
        checks = compare.compare_structure(original, recalculated)
        checks.append(compare.compare_formulas(original, recalculated))
        checks.append(compare.compare_inputs(original, recalculated, {}, self.tolerance))
        checks.append(compare.compare_errors(original, recalculated))
        cached = {(s, a): v for s, cells in original["cached_values"].items() for a, v in cells.items()}
        has_cache = any(v is not None for v in cached.values())
        if has_cache:
            checks.append(compare.compare_formula_values(cached, recalculated, self.tolerance,
                                                         code="matches_excel_cached_results",
                                                         label="the results stored by Excel"))
        if reference_values:
            checks.append(compare.compare_formula_values(reference_values, recalculated, self.tolerance,
                                                         code="matches_trusted_reference_outputs",
                                                         label="the trusted reference outputs"))
        passed = all(c.passed for c in checks)
        verified = has_cache or bool(reference_values)
        return {"passed": passed, "fidelity_verified": verified and passed,
                "checks": [c.as_dict() for c in checks]}
