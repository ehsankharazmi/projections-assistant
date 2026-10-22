"""Calculation worker: recalculates workbooks with headless LibreOffice Calc.

Run as `python -m app.calc.worker`. The worker only reads and writes the shared
jobs directory; it has no database access and needs no network.

For every job it uses a fresh LibreOffice profile configured to always
recalculate Excel files on load, to never execute macros, and to never update
external links; it separates input and output directories, enforces a timeout
(killing the whole process group), and verifies that the produced file is a
readable workbook whose formula cells all carry calculated results. A
successful process exit alone is never treated as success.
"""

from __future__ import annotations

import io
import json
import logging
import os
import shutil
import signal
import socket
import subprocess
import tempfile
import time
import zipfile
from pathlib import Path

from openpyxl import load_workbook

from app.calc.protocol import RecalcRequest, RecalcResult, job_dir, pending_jobs, try_claim, write_json_atomic

log = logging.getLogger("calc-worker")

EXPORT_FILTER = "Calc MS Excel 2007 XML"

# Calc/Formula/Load/OOXMLRecalcMode: 0 = always recalculate, 1 = never, 2 = prompt.
# Common/Security/Scripting/MacroSecurityLevel: 3 = very high (only trusted locations; none configured).
# Calc/Content/Update/Link: 0 = always, 1 = never, 2 = on request (no UI in headless mode).
PROFILE_SETTINGS = """<?xml version="1.0" encoding="UTF-8"?>
<oor:items xmlns:oor="http://openoffice.org/2001/registry" xmlns:xs="http://www.w3.org/2001/XMLSchema"
           xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">
<item oor:path="/org.openoffice.Office.Calc/Formula/Load"><prop oor:name="OOXMLRecalcMode" oor:op="fuse"><value>0</value></prop></item>
<item oor:path="/org.openoffice.Office.Calc/Formula/Load"><prop oor:name="ODFRecalcMode" oor:op="fuse"><value>0</value></prop></item>
<item oor:path="/org.openoffice.Office.Common/Security/Scripting"><prop oor:name="MacroSecurityLevel" oor:op="fuse"><value>3</value></prop></item>
<item oor:path="/org.openoffice.Office.Common/Security/Scripting"><prop oor:name="DisableMacrosExecution" oor:op="fuse"><value>true</value></prop></item>
<item oor:path="/org.openoffice.Office.Calc/Content/Update"><prop oor:name="Link" oor:op="fuse"><value>1</value></prop></item>
<item oor:path="/org.openoffice.Office.Common/Misc"><prop oor:name="UseLocking" oor:op="fuse"><value>false</value></prop></item>
</oor:items>
"""


class RecalcError(Exception):
    def __init__(self, code: str, message: str, diagnostics: dict | None = None):
        super().__init__(message)
        self.code = code
        self.message = message
        self.diagnostics = diagnostics or {}


def _write_profile(profile_dir: Path) -> None:
    user = profile_dir / "user"
    user.mkdir(parents=True, exist_ok=True)
    (user / "registrymodifications.xcu").write_text(PROFILE_SETTINGS)


def verify_output(data: bytes) -> dict:
    """Check that LibreOffice produced a readable workbook with calculated formula results."""
    try:
        zipfile.ZipFile(io.BytesIO(data)).testzip()
        formulas = load_workbook(io.BytesIO(data), data_only=False)
        values = load_workbook(io.BytesIO(data), data_only=True)
    except Exception as exc:
        raise RecalcError("invalid_output", "The calculation engine produced an unreadable workbook.",
                          {"exception": type(exc).__name__}) from exc
    formula_cells = missing = 0
    for ws in formulas.worksheets:
        wsv = values[ws.title]
        for row in ws.iter_rows():
            for cell in row:
                if cell.data_type == "f":
                    formula_cells += 1
                    if wsv[cell.coordinate].value is None:
                        missing += 1
    if missing:
        raise RecalcError("recalc_incomplete", "Some formulas were not calculated.",
                          {"formula_cells": formula_cells, "without_result": missing})
    return {"formula_cells": formula_cells}


def recalculate_bytes(data: bytes, *, soffice: str, timeout: float, work_root: Path | None = None) -> bytes:
    """Recalculate an .xlsx workbook with LibreOffice and return the saved result."""
    work = Path(tempfile.mkdtemp(prefix="recalc-", dir=work_root))
    try:
        in_dir, out_dir, profile = work / "in", work / "out", work / "profile"
        in_dir.mkdir()
        out_dir.mkdir()
        _write_profile(profile)
        source = in_dir / "workbook.xlsx"
        source.write_bytes(data)
        cmd = [soffice, "--headless", "--norestore", "--nologo", "--nodefault", "--nolockcheck",
               "--nofirststartwizard", f"-env:UserInstallation={profile.resolve().as_uri()}",
               "--convert-to", f"xlsx:{EXPORT_FILTER}", "--outdir", str(out_dir), str(source)]
        env = {"PATH": os.environ.get("PATH", "/usr/bin:/bin"), "HOME": str(work), "LANG": "C.UTF-8",
               "SAL_USE_VCLPLUGIN": "svp"}
        try:
            proc = subprocess.Popen(cmd, stdout=subprocess.PIPE, stderr=subprocess.PIPE, env=env,
                                    start_new_session=True, cwd=work)
        except FileNotFoundError as exc:
            raise RecalcError("engine_unavailable", "LibreOffice is not installed on the calculation worker.") from exc
        try:
            stdout, stderr = proc.communicate(timeout=timeout)
        except subprocess.TimeoutExpired:
            os.killpg(proc.pid, signal.SIGKILL)
            proc.communicate()
            raise RecalcError("engine_timeout", "The calculation took too long and was stopped.",
                              {"timeout_seconds": timeout}) from None
        output = out_dir / "workbook.xlsx"
        if proc.returncode != 0 or not output.exists():
            raise RecalcError("engine_failed", "The calculation engine could not process the workbook.",
                              {"returncode": proc.returncode, "stderr": stderr.decode(errors="replace")[-500:]})
        result = output.read_bytes()
        verify_output(result)
        return result
    finally:
        shutil.rmtree(work, ignore_errors=True)


def process_job(jobs_dir: Path, job_id: str, *, soffice: str) -> RecalcResult:
    directory = job_dir(jobs_dir, job_id)
    request = RecalcRequest(**json.loads((directory / "request.json").read_text()))
    started = time.monotonic()
    try:
        output = recalculate_bytes((directory / "input.xlsx").read_bytes(), soffice=soffice,
                                   timeout=request.timeout_seconds, work_root=directory)
        tmp = directory / ".output.tmp"
        tmp.write_bytes(output)
        os.replace(tmp, directory / "output.xlsx")
        result = RecalcResult(job_id=job_id, ok=True, duration_seconds=time.monotonic() - started)
    except RecalcError as exc:
        result = RecalcResult(job_id=job_id, ok=False, error_code=exc.code, error_message=exc.message,
                              duration_seconds=time.monotonic() - started, diagnostics=exc.diagnostics)
    except Exception as exc:  # never leave a job without a result
        log.exception("unexpected worker failure for %s", job_id)
        result = RecalcResult(job_id=job_id, ok=False, error_code="worker_error",
                              error_message="The calculation worker failed unexpectedly.",
                              duration_seconds=time.monotonic() - started,
                              diagnostics={"exception": type(exc).__name__})
    write_json_atomic(directory / "result.json", result.__dict__)
    return result


def recover_interrupted(jobs_dir: Path) -> None:
    """On startup, release jobs this (single) worker claimed but never finished, allowing one retry."""
    if not jobs_dir.exists():
        return
    for entry in jobs_dir.iterdir():
        claim = entry / "claim.json"
        if claim.exists() and not (entry / "result.json").exists():
            request_path = entry / "request.json"
            request = json.loads(request_path.read_text())
            if request.get("attempt", 1) >= 2:
                write_json_atomic(entry / "result.json", RecalcResult(
                    job_id=entry.name, ok=False, error_code="interrupted",
                    error_message="The calculation was interrupted twice and was not retried again.").__dict__)
            else:
                request["attempt"] = request.get("attempt", 1) + 1
                write_json_atomic(request_path, request)
                claim.unlink()


def run_forever(jobs_dir: Path, *, soffice: str, poll_seconds: float = 0.5, stop=None) -> None:
    worker_id = f"{socket.gethostname()}-{os.getpid()}"
    jobs_dir.mkdir(parents=True, exist_ok=True)
    recover_interrupted(jobs_dir)
    log.info("calculation worker %s watching %s", worker_id, jobs_dir)
    heartbeat = jobs_dir / ".worker-heartbeat"
    while stop is None or not stop.is_set():
        try:
            heartbeat.touch()
            for job_id in pending_jobs(jobs_dir):
                if try_claim(jobs_dir, job_id, worker_id):
                    result = process_job(jobs_dir, job_id, soffice=soffice)
                    log.info("job %s finished ok=%s code=%s in %.1fs", job_id, result.ok, result.error_code,
                             result.duration_seconds)
        except FileNotFoundError:
            pass  # a job directory disappeared between listing and claiming
        except Exception:
            log.exception("worker iteration failed")
        time.sleep(poll_seconds)


def main() -> None:  # pragma: no cover - process entry point
    logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(name)s %(message)s")
    jobs_dir = Path(os.environ.get("JOBS_DIR", "data/jobs"))
    run_forever(jobs_dir, soffice=os.environ.get("SOFFICE_PATH", "soffice"))


if __name__ == "__main__":
    main()
