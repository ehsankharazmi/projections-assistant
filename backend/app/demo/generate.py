"""Generate the synthetic demo workbook, its trusted reference outputs, and its demo map.

The workbook is written with openpyxl, then the cached formula values that Excel
would normally store are injected from `model.compute()`. That makes the demo
file behave like a workbook last saved by Excel, so the compatibility gate can
compare LibreOffice's recalculation against stored values.

Usage: python -m app.demo.generate <output-dir>
"""

from __future__ import annotations

import datetime as dt
import io
import json
import sys
import zipfile
from pathlib import Path

from lxml import etree
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.workbook.defined_name import DefinedName

from app.demo.model import BASE_INPUTS, PRODUCTS, compute, scenario_inputs

WORKBOOK_FILENAME = "demo_product_margins.xlsx"
REFERENCE_FILENAME = "demo_reference_outputs.json"
MAP_FILENAME = "demo_map.json"

MONEY = "#,##0.00"
WHOLE = "#,##0"
PCT = "0.0%"
INPUT_FILL = PatternFill("solid", fgColor="FFF4CC")
FIRST_ROW = 4
LAST_ROW = FIRST_ROW + len(PRODUCTS) - 1
TOTAL_ROW = LAST_ROW + 1

# Products sheet output columns: (column, field id, label, unit, value_type, number format)
PRODUCT_COLUMNS = [
    ("G", "revenue", "Revenue", "USD/year", "currency", MONEY),
    ("H", "manufacturing_labor", "Manufacturing labor cost", "USD/year", "currency", MONEY),
    ("I", "packaging_labor", "Packaging labor cost", "USD/year", "currency", MONEY),
    ("J", "materials", "Materials cost", "USD/year", "currency", MONEY),
    ("K", "manufacturing_hours", "Manufacturing hours", "hours/year", "number", WHOLE),
    ("L", "overhead", "Allocated overhead", "USD/year", "currency", MONEY),
    ("M", "total_cost", "Total cost", "USD/year", "currency", MONEY),
    ("N", "profit", "Profit", "USD/year", "currency", MONEY),
    ("O", "margin", "Profit margin", "fraction", "percent", PCT),
    ("P", "loss_flag", "Loses money (1 = yes)", "flag", "number", WHOLE),
]

SUMMARY_ROWS = [
    (4, "total_revenue", "Total revenue", "=Products!G{t}", "USD/year", "currency", MONEY),
    (5, "total_cost", "Total cost", "=Products!M{t}", "USD/year", "currency", MONEY),
    (6, "total_profit", "Total profit", "=Products!N{t}", "USD/year", "currency", MONEY),
    (7, "overall_margin", "Overall profit margin", "=Products!O{t}", "fraction", "percent", PCT),
    (8, "products_losing_money", "Number of products losing money", "=Products!P{t}", "count", "number", WHOLE),
]


def _product_formulas(r: int) -> dict[str, str]:
    return {
        "G": f"=B{r}*C{r}",
        "H": f"=C{r}*D{r}*Inputs!$B$4*(1+Inputs!$B$6)",
        "I": f"=C{r}*E{r}*Inputs!$B$5*(1+Inputs!$B$6)",
        "J": f"=C{r}*F{r}",
        "K": f"=C{r}*D{r}",
        "L": f"=Inputs!$B$7*K{r}/SUM($K${FIRST_ROW}:$K${LAST_ROW})",
        "M": f"=H{r}+I{r}+J{r}+L{r}",
        "N": f"=G{r}-M{r}",
        "O": f"=IF(G{r}=0,0,N{r}/G{r})",
        "P": f"=IF(N{r}<0,1,0)",
    }


def build_workbook() -> tuple[Workbook, dict[tuple[str, str], float]]:
    """Return the workbook and the cached value for every formula cell."""
    results = compute()
    cached: dict[tuple[str, str], float] = {}
    wb = Workbook()
    bold = Font(bold=True)

    ws = wb.active
    ws.title = "Inputs"
    ws["A1"] = "Synthetic demo workbook - all products and figures are invented"
    ws["A1"].font = bold
    for col, text in zip("ABCD", ["Input", "Value", "Unit", "Notes"], strict=True):
        ws[f"{col}3"] = text
        ws[f"{col}3"].font = bold
    inputs = [
        (4, "Manufacturing hourly wage", BASE_INPUTS.manufacturing_hourly_wage, "USD/hour", MONEY,
         "Base hourly wage for manufacturing labor"),
        (5, "Packaging hourly wage", BASE_INPUTS.packaging_hourly_wage, "USD/hour", MONEY,
         "Base hourly wage for packaging labor"),
        (6, "Payroll burden rate", BASE_INPUTS.payroll_burden_rate, "fraction", PCT,
         "Benefits and payroll taxes as a share of wages"),
        (7, "Annual fixed overhead", BASE_INPUTS.annual_fixed_overhead, "USD/year", MONEY,
         "Allocated to products by share of manufacturing hours"),
    ]
    for row, label, value, unit, fmt, note in inputs:
        ws[f"A{row}"] = label
        ws[f"B{row}"] = value
        ws[f"B{row}"].number_format = fmt
        ws[f"B{row}"].fill = INPUT_FILL
        ws[f"C{row}"] = unit
        ws[f"D{row}"] = note
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["D"].width = 52

    ws = wb.create_sheet("Products")
    ws["A1"] = "Product economics (invented data)"
    ws["A1"].font = bold
    headers = ["Product", "Price", "Annual volume", "Manufacturing hours per unit",
               "Packaging hours per unit", "Material cost per unit"] + [c[2] for c in PRODUCT_COLUMNS]
    for idx, text in enumerate(headers):
        cell = ws.cell(row=3, column=idx + 1, value=text)
        cell.font = bold
    for i, p in enumerate(PRODUCTS):
        r = FIRST_ROW + i
        ws[f"A{r}"] = p.name
        for col, value, fmt in (("B", p.price, MONEY), ("C", p.volume, WHOLE), ("D", p.mfg_hours, "0.00"),
                                ("E", p.pkg_hours, "0.00"), ("F", p.material_cost, MONEY)):
            ws[f"{col}{r}"] = value
            ws[f"{col}{r}"].number_format = fmt
            ws[f"{col}{r}"].fill = INPUT_FILL
        formulas = _product_formulas(r)
        row_result = results["products"][i]
        for col, field, _label, _unit, _vt, fmt in PRODUCT_COLUMNS:
            ws[f"{col}{r}"] = formulas[col]
            ws[f"{col}{r}"].number_format = fmt
            cached[("Products", f"{col}{r}")] = row_result[field]
    t = TOTAL_ROW
    ws[f"A{t}"] = "Total"
    ws[f"A{t}"].font = bold
    for col, field, _label, _unit, _vt, fmt in PRODUCT_COLUMNS:
        if col == "O":
            ws[f"O{t}"] = f"=IF(G{t}=0,0,N{t}/G{t})"
            cached[("Products", f"O{t}")] = results["summary"]["overall_margin"]
        else:
            ws[f"{col}{t}"] = f"=SUM({col}{FIRST_ROW}:{col}{LAST_ROW})"
            cached[("Products", f"{col}{t}")] = sum(row[field] for row in results["products"])
        ws[f"{col}{t}"].number_format = fmt
        ws[f"{col}{t}"].font = bold
    ws.column_dimensions["A"].width = 24

    ws = wb.create_sheet("Summary")
    ws["A1"] = "Portfolio summary (invented data)"
    ws["A1"].font = bold
    ws["A3"], ws["B3"] = "Measure", "Value"
    ws["A3"].font = ws["B3"].font = bold
    for row, field, label, formula, _unit, _vt, fmt in SUMMARY_ROWS:
        ws[f"A{row}"] = label
        ws[f"B{row}"] = formula.format(t=TOTAL_ROW)
        ws[f"B{row}"].number_format = fmt
        cached[("Summary", f"B{row}")] = results["summary"][field]
    ws.column_dimensions["A"].width = 34

    wb.defined_names["ManufacturingHourlyWage"] = DefinedName("ManufacturingHourlyWage", attr_text="Inputs!$B$4")
    wb.defined_names["AnnualFixedOverhead"] = DefinedName("AnnualFixedOverhead", attr_text="Inputs!$B$7")
    fixed = dt.datetime(2026, 1, 1, 0, 0, 0)
    wb.properties.creator = "projections-assistant demo generator"
    wb.properties.title = "Synthetic product-margin demo (invented data)"
    wb.properties.created = fixed
    wb.properties.modified = fixed
    return wb, cached


def _format_number(value: float) -> str:
    if float(value).is_integer():
        return str(int(value))
    return repr(float(value))


def inject_cached_values(xlsx_bytes: bytes, cached: dict[tuple[str, str], float]) -> bytes:
    """Add <v> cached values to formula cells, as Excel does when it saves."""
    ns = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
          "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships"}
    src = zipfile.ZipFile(io.BytesIO(xlsx_bytes))
    workbook_xml = etree.fromstring(src.read("xl/workbook.xml"))
    rels_xml = etree.fromstring(src.read("xl/_rels/workbook.xml.rels"))
    targets = {rel.get("Id"): rel.get("Target") for rel in rels_xml}
    sheet_paths: dict[str, str] = {}
    for sheet in workbook_xml.find("m:sheets", ns):
        rid = sheet.get(f"{{{ns['r']}}}id")
        target = targets[rid].lstrip("/")
        sheet_paths[target if target.startswith("xl/") else f"xl/{target}"] = sheet.get("name")

    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            data = src.read(info.filename)
            sheet_name = sheet_paths.get(info.filename)
            if sheet_name is not None:
                root = etree.fromstring(data)
                for cell in root.iter(f"{{{ns['m']}}}c"):
                    formula = cell.find("m:f", ns)
                    if formula is None:
                        continue
                    key = (sheet_name, cell.get("r"))
                    if key not in cached:
                        raise KeyError(f"No cached value computed for {key}")
                    for old in cell.findall("m:v", ns):
                        cell.remove(old)
                    value = etree.SubElement(cell, f"{{{ns['m']}}}v")
                    value.text = _format_number(cached[key])
                data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            fixed_info = zipfile.ZipInfo(info.filename, date_time=(2026, 1, 1, 0, 0, 0))
            fixed_info.compress_type = zipfile.ZIP_DEFLATED
            dst.writestr(fixed_info, data)
    return out.getvalue()


def reference_cells(results: dict) -> dict[str, float]:
    """Express model results as {"Sheet!A1": value} for every formula cell in the workbook."""
    cells: dict[str, float] = {}
    for i, row in enumerate(results["products"]):
        for col, field, *_ in PRODUCT_COLUMNS:
            cells[f"Products!{col}{FIRST_ROW + i}"] = row[field]
    for col, field, *_ in PRODUCT_COLUMNS:
        cells[f"Products!{col}{TOTAL_ROW}"] = (results["summary"]["overall_margin"] if col == "O"
                                               else sum(r[field] for r in results["products"]))
    for row, field, *_ in SUMMARY_ROWS:
        cells[f"Summary!B{row}"] = results["summary"][field]
    return cells


def reference_outputs() -> dict:
    """Trusted expected outputs computed independently of the workbook formulas."""
    scenarios = {"baseline": 1.0, "manufacturing_wage_plus_5pct": 1.05, "manufacturing_wage_plus_8pct": 1.08}
    computed = {name: compute(scenario_inputs(factor)) for name, factor in scenarios.items()}
    return {
        "description": "Trusted reference outputs for the synthetic demo workbook, computed in plain Python "
                       "(app.demo.model.compute), independent of spreadsheet formula evaluation.",
        "baseline_cells": reference_cells(computed["baseline"]),
        "scenarios": {
            name: {"manufacturing_wage_factor": factor, "results": computed[name],
                   "cells": reference_cells(computed[name])}
            for name, factor in scenarios.items()
        },
    }


def demo_map() -> dict:
    """Mapping for the demo workbook. Marked source=demo; confirmed because it ships with the demo."""
    fields = [
        {"id": "manufacturing_hourly_rate", "label": "Manufacturing hourly wage",
         "description": "Base hourly wage used for manufacturing labor costs (before payroll burden)",
         "sheet": "Inputs", "address": "B4", "role": "input", "value_type": "currency", "unit": "USD/hour",
         "minimum": 0, "maximum": 500, "aliases": ["manufacturing wages", "hourly labor rate", "factory wage"],
         "editable": True},
        {"id": "packaging_hourly_rate", "label": "Packaging hourly wage",
         "description": "Base hourly wage used for packaging labor costs (before payroll burden)",
         "sheet": "Inputs", "address": "B5", "role": "input", "value_type": "currency", "unit": "USD/hour",
         "minimum": 0, "maximum": 500, "aliases": ["packaging wages"], "editable": True},
        {"id": "payroll_burden_rate", "label": "Payroll burden rate",
         "description": "Benefits and payroll taxes as a share of wages; applies to manufacturing and packaging",
         "sheet": "Inputs", "address": "B6", "role": "input", "value_type": "percent", "unit": "fraction",
         "minimum": 0, "maximum": 1, "aliases": ["benefits", "payroll taxes", "labor burden"], "editable": True},
        {"id": "annual_fixed_overhead", "label": "Annual fixed overhead",
         "description": "Fixed overhead allocated to products by share of manufacturing hours",
         "sheet": "Inputs", "address": "B7", "role": "input", "value_type": "currency", "unit": "USD/year",
         "minimum": 0, "maximum": None, "aliases": ["overhead", "fixed costs"], "editable": True},
    ]
    for row, field, label, _formula, unit, vt, _fmt in SUMMARY_ROWS:
        fields.append({"id": field, "label": label, "description": f"{label} across all products",
                       "sheet": "Summary", "address": f"B{row}", "role": "output", "value_type": vt,
                       "unit": unit, "minimum": None, "maximum": None, "aliases": [], "editable": False})
    for f in fields:
        f["confirmed"] = True
        f["source"] = "demo"
    table = {
        "id": "product_economics", "label": "Product economics", "sheet": "Products",
        "description": "Per-product revenue, costs, profit, and margin",
        "row_label_column": "A", "first_row": FIRST_ROW, "last_row": LAST_ROW, "total_row": TOTAL_ROW,
        "columns": [{"id": field, "label": label, "column": col, "unit": unit, "value_type": vt}
                    for col, field, label, unit, vt, _fmt in PRODUCT_COLUMNS],
        "confirmed": True,
    }
    return {
        "schema_version": 1,
        "source": "demo",
        "fields": fields,
        "output_tables": [table],
        "output_field_ids": [row[1] for row in SUMMARY_ROWS],
        "business_rules": [
            {"id": "overhead_allocation",
             "text": "Annual fixed overhead is allocated to products by each product's share of total manufacturing hours."},
            {"id": "labor_cost_definition",
             "text": "Labor cost = hours x hourly wage x (1 + payroll burden rate). 'Labor costs' can mean wages, "
                     "hours, or burden, so ask which one before changing anything."},
            {"id": "percent_storage",
             "text": "Percentages are stored as fractions (0.20 = 20%). A '% increase' to a wage is relative; "
                     "a change to a rate given in 'points' is absolute."},
        ],
    }


def generate(out_dir: Path) -> dict[str, Path]:
    out_dir.mkdir(parents=True, exist_ok=True)
    wb, cached = build_workbook()
    buf = io.BytesIO()
    wb.save(buf)
    paths = {
        "workbook": out_dir / WORKBOOK_FILENAME,
        "reference": out_dir / REFERENCE_FILENAME,
        "map": out_dir / MAP_FILENAME,
    }
    paths["workbook"].write_bytes(inject_cached_values(buf.getvalue(), cached))
    paths["reference"].write_text(json.dumps(reference_outputs(), indent=2))
    paths["map"].write_text(json.dumps(demo_map(), indent=2))
    return paths


if __name__ == "__main__":
    target = Path(sys.argv[1]) if len(sys.argv) > 1 else Path("demo/generated")
    for kind, path in generate(target).items():
        print(f"{kind}: {path}")
