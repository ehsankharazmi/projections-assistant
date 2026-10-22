"""Acceptance checks 7 and 8: honest handling of calculation problems; the original is always retained."""

from __future__ import annotations

import io

from openpyxl import Workbook

import app.calc.worker as worker
import app.services.jobs as jobs_module
from app.calc.worker import RecalcError
from tests.conftest import H, fake_recalculate, load_demo, run_scenario, wait_job


def _upload(client, wb: Workbook, name="custom.xlsx") -> dict:
    buf = io.BytesIO()
    wb.save(buf)
    r = client.post("/api/workbooks", headers=H, files={"file": (name, buf.getvalue(), "application/octet-stream")})
    assert r.status_code == 201, r.text
    return r.json()


def test_missing_caches_require_baseline_confirmation(make_client, monkeypatch):
    def recalc_simple(data, **kw):
        from openpyxl import load_workbook

        from app.demo.generate import inject_cached_values
        wb = load_workbook(io.BytesIO(data))
        out = io.BytesIO()
        wb.save(out)
        return inject_cached_values(out.getvalue(), {("Sheet", "B1"): 2.0})

    client = make_client()
    monkeypatch.setattr(worker, "recalculate_bytes", recalc_simple)
    wb = Workbook()
    wb.active.title = "Sheet"
    wb.active["A1"] = 1
    wb.active["B1"] = "=SUM(A1,1)"  # openpyxl writes no cached result
    body = _upload(client, wb)
    assert wait_job(client, body["job_id"])["status"] == "succeeded"
    info = client.get(f"/api/workbooks/{body['workbook']['id']}").json()
    assert info["compatibility_status"] == "needs_baseline_confirmation"
    assert info["map"]["fields"] == [] or all(not f["editable"] for f in info["map"]["fields"])
    r = client.post(f"/api/workbooks/{body['workbook']['id']}/confirm-baseline", headers=H)
    assert r.status_code == 200 and r.json()["compatibility_status"] == "editable"


def test_reference_outputs_can_verify_or_block(make_client, monkeypatch):
    def recalc_simple(data, **kw):
        from openpyxl import load_workbook

        from app.demo.generate import inject_cached_values
        out = io.BytesIO()
        load_workbook(io.BytesIO(data)).save(out)
        return inject_cached_values(out.getvalue(), {("Sheet", "B1"): 2.0})

    client = make_client()
    monkeypatch.setattr(worker, "recalculate_bytes", recalc_simple)
    wb = Workbook()
    wb.active.title = "Sheet"
    wb.active["A1"] = 1
    wb.active["B1"] = "=SUM(A1,1)"
    body = _upload(client, wb)
    wait_job(client, body["job_id"])
    wid = body["workbook"]["id"]
    wrong = client.put(f"/api/workbooks/{wid}/reference-outputs", headers=H, json={"cells": {"Sheet!B1": 3}})
    assert wrong.json()["check"]["passed"] is False and wrong.json()["compatibility_status"] == "inspect_only"


def test_baseline_mismatch_blocks_editing(make_client, monkeypatch):
    def wrong_engine(data, **kw):
        out = fake_recalculate(data, **kw)
        from openpyxl import load_workbook

        from app.demo.generate import inject_cached_values, reference_cells
        from app.demo.model import compute
        cells = reference_cells(compute())
        cells["Summary!B6"] += 1000  # engine disagrees with Excel's stored total profit
        buf = io.BytesIO()
        load_workbook(io.BytesIO(out)).save(buf)
        return inject_cached_values(buf.getvalue(), {tuple(k.split("!", 1)): v for k, v in cells.items()})

    client = make_client()
    monkeypatch.setattr(worker, "recalculate_bytes", wrong_engine)
    body = client.post("/api/workbooks/demo", headers=H).json()
    wait_job(client, body["job_id"])
    info = client.get(f"/api/workbooks/{body['workbook']['id']}").json()
    assert info["compatibility_status"] == "inspect_only"
    assert any("differ" in m for m in info["compatibility"]["round_trip"]["failed_checks"])
    r = client.post(f"/api/workbooks/{info['id']}/scenarios", headers=H, json={
        "base_version_id": info["active_version"]["id"], "name": "x",
        "changes": [{"field_id": "manufacturing_hourly_rate", "operation": "multiply", "operand": 1.08}]})
    assert r.status_code == 422 and r.json()["error"]["code"] == "editing_blocked"


def test_unsupported_features_are_inspect_only(make_client):
    client = make_client()
    wb = Workbook()
    wb.active["A1"] = "Rate"
    wb.active["B1"] = 3
    wb.active["C1"] = "=NOW()"
    body = _upload(client, wb)
    assert body["job_id"] is None  # no round trip attempted when preflight blocks
    info = body["workbook"]
    assert info["compatibility_status"] == "inspect_only"
    assert info["compatibility"]["blocking"][0]["code"] == "volatile_functions"


def test_engine_timeout_fails_scenario_and_keeps_original(make_client, monkeypatch):
    client = make_client()
    wb = load_demo(client)

    def timeout(data, **kw):
        raise RecalcError("engine_timeout", "The calculation took too long and was stopped.")

    monkeypatch.setattr(worker, "recalculate_bytes", timeout)
    scenario = run_scenario(client, wb, 1.08)
    assert scenario["status"] == "failed" and scenario["error"]["code"] == "engine_timeout"
    assert scenario["apply_allowed"] is False
    r = client.post(f"/api/scenarios/{scenario['id']}/apply", headers={**H, "Idempotency-Key": "t"},
                    json={"expected_active_version_id": wb["active_version"]["id"]})
    assert r.status_code == 409
    assert client.get(f"/api/workbooks/{wb['id']}").json()["active_version"]["id"] == wb["active_version"]["id"]


def test_new_formula_errors_block_apply(make_client, monkeypatch):
    client = make_client()
    wb = load_demo(client)

    def with_error(data, **kw):
        out = fake_recalculate(data, **kw)
        import zipfile

        from lxml import etree
        src = zipfile.ZipFile(io.BytesIO(out))
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w") as dst:
            for info in src.infolist():
                content = src.read(info.filename)
                if info.filename == "xl/worksheets/sheet2.xml":
                    root = etree.fromstring(content)
                    ns = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
                    cell = root.find(".//m:c[@r='O5']", ns)
                    cell.set("t", "e")
                    cell.find("m:v", ns).text = "#DIV/0!"
                    content = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
                dst.writestr(info, content)
        return buf.getvalue()

    monkeypatch.setattr(worker, "recalculate_bytes", with_error)
    scenario = run_scenario(client, wb, 1.08)
    assert scenario["status"] == "failed" and scenario["error"]["code"] == "validation_failed"
    failed = [v for v in scenario["validations"] if not v["passed"]]
    assert failed[0]["code"] == "no_new_formula_errors"
    assert failed[0]["details"]["new_errors"][0]["cell"] == "Products!O5"


def test_unavailable_calculation_service_reported(make_client, monkeypatch):
    monkeypatch.setattr(jobs_module, "UNCLAIMED_GRACE_SECONDS", 1.0)
    client = make_client(worker_mode="external")  # no worker running at all
    body = client.post("/api/workbooks/demo", headers=H).json()
    job = wait_job(client, body["job_id"], timeout=30)
    assert job["status"] == "failed" and job["error"]["code"] == "calculation_service_unavailable"
    info = client.get(f"/api/workbooks/{body['workbook']['id']}").json()
    assert info["compatibility_status"] == "inspect_only"
    health = client.get("/api/health").json()
    assert health["components"]["calculation_worker"] == "unavailable" and health["status"] == "degraded"
    # Deterministic operations still work without the calculation service.
    export = client.get(f"/api/workbooks/{info['id']}/export", params={"version_id": info["active_version"]["id"]})
    assert export.status_code == 200
