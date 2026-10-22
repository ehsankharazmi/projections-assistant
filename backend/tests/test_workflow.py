"""End-to-end workflow through the REST API (calculation by the demo test double).

Covers acceptance checks 1, 3, 5, 9 (stale previews, repeated Apply, mapping
changes, simultaneous tabs), and 12 (deletion).
"""

from __future__ import annotations

import io
import json

import pytest
from openpyxl import load_workbook

from app.demo.model import compute, scenario_inputs
from tests.conftest import H, load_demo, run_scenario


def _profit_by_product(comparison: dict) -> dict[str, float]:
    table = comparison["tables"][0]
    return {r["label"]: r["cells"]["profit"]["after"] for r in table["rows"]}


def _expected_profit(factor: float) -> dict[str, float]:
    return {r["product"]: r["profit"] for r in compute(scenario_inputs(factor))["products"]}


def test_demo_import_passes_compatibility_gate(make_client):
    client = make_client()
    wb = load_demo(client)
    assert wb["compatibility_status"] == "editable"
    assert wb["compatibility"]["round_trip"]["passed"] is True
    assert wb["compatibility"]["round_trip"]["fidelity_verified"] is True
    assert wb["active_version"]["seq"] == 1
    assert wb["is_demo"] is True


def test_wage_increase_preview_apply_export(make_client):
    client = make_client()
    wb = load_demo(client)
    scenario = run_scenario(client, wb, 1.08)
    assert scenario["status"] == "ready", scenario
    assert scenario["apply_allowed"] is True
    assert all(v["passed"] for v in scenario["validations"])
    change = scenario["changes"][0]
    assert change["old_value"] == 32.0 and change["new_value"] == pytest.approx(34.56)
    expected = _expected_profit(1.08)
    for product, profit in _profit_by_product(scenario["comparison"]).items():
        assert profit == pytest.approx(expected[product], rel=1e-9)
    flips = scenario["comparison"]["tables"][0]["sign_changes"]
    assert {(f["row_label"], f["column_id"], f["change"]) for f in flips} >= {
        ("Walnut Side Table", "profit", "became_negative")}
    # A preview changes neither the active version nor the original.
    assert client.get(f"/api/workbooks/{wb['id']}").json()["active_version"]["id"] == wb["active_version"]["id"]

    r = client.post(f"/api/scenarios/{scenario['id']}/apply", headers={**H, "Idempotency-Key": "k-1"},
                    json={"expected_active_version_id": wb["active_version"]["id"]})
    assert r.status_code == 200, r.text
    new_version = r.json()["version"]
    assert new_version["seq"] == 2 and new_version["parent_version_id"] == wb["active_version"]["id"]
    after = client.get(f"/api/workbooks/{wb['id']}").json()
    assert after["active_version"]["id"] == new_version["id"]

    export = client.get(f"/api/workbooks/{wb['id']}/export", params={"version_id": new_version["id"]})
    assert export.status_code == 200
    assert export.headers["x-version-id"] == new_version["id"]
    exported = load_workbook(io.BytesIO(export.content))
    assert exported["Inputs"]["B4"].value == pytest.approx(34.56)
    assert exported["Products"]["N5"].value == "=G5-M5"  # formulas intact
    assert exported["Products"]["L4"].value.startswith("=Inputs!$B$7*K4/SUM(")
    original = client.get(f"/api/workbooks/{wb['id']}/export", params={"version_id": wb["active_version"]["id"]})
    assert load_workbook(io.BytesIO(original.content))["Inputs"]["B4"].value == 32.0  # original retained


def test_revision_replaces_preview_without_compounding(make_client):
    client = make_client()
    wb = load_demo(client)
    first = run_scenario(client, wb, 1.08)
    revised = run_scenario(client, wb, 1.05, revision_of=first["id"])
    assert revised["base_version"]["id"] == first["base_version"]["id"]
    assert revised["changes"][0]["old_value"] == 32.0
    assert revised["changes"][0]["new_value"] == pytest.approx(33.6)  # 32 x 1.05, not 32 x 1.08 x 1.05
    expected = _expected_profit(1.05)
    for product, profit in _profit_by_product(revised["comparison"]).items():
        assert profit == pytest.approx(expected[product], rel=1e-9)
    assert expected["Walnut Side Table"] > 0  # still profitable at 5%
    old = client.get(f"/api/scenarios/{first['id']}").json()
    assert old["status"] == "superseded"


def test_repeated_apply_is_idempotent_and_stale_preview_rejected(make_client):
    client = make_client()
    wb = load_demo(client)
    a = run_scenario(client, wb, 1.08)
    b = run_scenario(client, wb, 1.05)  # a second, independent preview on the same base (another tab)
    body = {"expected_active_version_id": wb["active_version"]["id"]}
    first = client.post(f"/api/scenarios/{a['id']}/apply", headers={**H, "Idempotency-Key": "same"}, json=body)
    again = client.post(f"/api/scenarios/{a['id']}/apply", headers={**H, "Idempotency-Key": "same"}, json=body)
    assert first.status_code == again.status_code == 200
    assert first.json() == again.json()
    versions = client.get(f"/api/workbooks/{wb['id']}/versions").json()
    assert len(versions) == 2  # no duplicate version
    new_key = client.post(f"/api/scenarios/{a['id']}/apply", headers={**H, "Idempotency-Key": "other"}, json=body)
    assert new_key.status_code == 409 and new_key.json()["error"]["code"] == "already_applied"
    stale = client.post(f"/api/scenarios/{b['id']}/apply", headers={**H, "Idempotency-Key": "b"}, json=body)
    assert stale.status_code == 409 and stale.json()["error"]["code"] == "stale_preview"
    assert client.get(f"/api/scenarios/{b['id']}").json()["apply_allowed"] is False
    missing_key = client.post(f"/api/scenarios/{b['id']}/apply", headers=H, json=body)
    assert missing_key.status_code == 400


def test_mapping_change_invalidates_pending_previews(make_client):
    client = make_client()
    wb = load_demo(client)
    scenario = run_scenario(client, wb, 1.08)
    new_map = {k: v for k, v in wb["map"].items() if k not in ("workbook_id", "mapping_version_id",
                                                                 "source_version_id")}
    new_map["fields"][0]["label"] = "Factory hourly wage"
    r = client.put(f"/api/workbooks/{wb['id']}/map", headers=H,
                   json={"expected_mapping_version_id": wb["mapping_version_id"], "map": new_map})
    assert r.status_code == 200, r.text
    assert client.get(f"/api/scenarios/{scenario['id']}").json()["status"] == "superseded"
    stale = client.put(f"/api/workbooks/{wb['id']}/map", headers=H,
                       json={"expected_mapping_version_id": wb["mapping_version_id"], "map": new_map})
    assert stale.status_code == 409


def test_restore_previous_version_restores_file_and_map(make_client):
    client = make_client()
    wb = load_demo(client)
    original_export = client.get(f"/api/workbooks/{wb['id']}/export",
                                 params={"version_id": wb["active_version"]["id"]}).content
    scenario = run_scenario(client, wb, 1.08)
    applied = client.post(f"/api/scenarios/{scenario['id']}/apply", headers={**H, "Idempotency-Key": "x"},
                          json={"expected_active_version_id": wb["active_version"]["id"]}).json()["version"]
    bad = client.post(f"/api/workbooks/{wb['id']}/restore", headers=H,
                      json={"target_version_id": wb["active_version"]["id"],
                            "expected_active_version_id": wb["active_version"]["id"]})
    assert bad.status_code == 409  # stale expectation
    r = client.post(f"/api/workbooks/{wb['id']}/restore", headers=H,
                    json={"target_version_id": wb["active_version"]["id"], "expected_active_version_id": applied["id"]})
    assert r.status_code == 200, r.text
    now = client.get(f"/api/workbooks/{wb['id']}").json()
    assert now["active_version"]["id"] == wb["active_version"]["id"]
    assert now["mapping_version_id"] == wb["active_version"]["mapping_version_id"]
    restored = client.get(f"/api/workbooks/{wb['id']}/export", params={"version_id": now["active_version"]["id"]})
    assert restored.content == original_export  # exact file identity
    assert restored.headers["x-version-sha256"] == wb["active_version"]["sha256"]


def test_delete_removes_records_and_files(make_client, tmp_path):
    client = make_client()
    wb = load_demo(client)
    run_scenario(client, wb, 1.08)
    workbook_dir = tmp_path / "data" / "workbooks" / wb["id"]
    assert workbook_dir.exists()
    assert client.delete(f"/api/workbooks/{wb['id']}", headers=H).status_code == 400  # needs confirmation
    r = client.delete(f"/api/workbooks/{wb['id']}", headers={**H, "X-Confirm-Delete": wb["id"]})
    assert r.status_code == 204
    assert not workbook_dir.exists()
    assert client.get(f"/api/workbooks/{wb['id']}").status_code == 404
    from sqlalchemy import select

    from app.db import session_scope
    from app.models import AuditEvent, Conversation, MappingVersion, Message, Scenario, WorkbookVersion
    with session_scope() as session:
        for model in (WorkbookVersion, MappingVersion, Scenario, Conversation):
            assert session.scalars(select(model).where(model.workbook_id == wb["id"])).first() is None
        assert session.scalars(select(Message)).first() is None
        assert session.scalars(select(AuditEvent).where(AuditEvent.workbook_id == wb["id"])).first() is None
    jobs_left = [p for p in (tmp_path / "data" / "jobs").iterdir() if p.name.startswith("job_")]
    assert jobs_left == []


def test_compare_versions_uses_explicit_base(make_client):
    client = make_client()
    wb = load_demo(client)
    scenario = run_scenario(client, wb, 1.08)
    applied = client.post(f"/api/scenarios/{scenario['id']}/apply", headers={**H, "Idempotency-Key": "c"},
                          json={"expected_active_version_id": wb["active_version"]["id"]}).json()["version"]
    cmp = client.get(f"/api/workbooks/{wb['id']}/compare", params={"base_version_id": wb["active_version"]["id"],
                                                                    "compared_version_id": applied["id"]}).json()
    assert cmp["base_version"]["id"] == wb["active_version"]["id"]
    profit = next(o for o in cmp["outputs"] if o["field_id"] == "total_profit")
    base = compute(scenario_inputs(1.0))["summary"]["total_profit"]
    new = compute(scenario_inputs(1.08))["summary"]["total_profit"]
    assert profit["delta"] == pytest.approx(new - base)
    assert profit["percent_change"] == pytest.approx((new - base) / base)
    margin = next(o for o in cmp["outputs"] if o["field_id"] == "overall_margin")
    assert margin["percentage_point_change"] == pytest.approx(margin["delta"] * 100)


def test_cross_user_access_is_not_found(make_client):
    client = make_client()
    wb = load_demo(client)
    from app.db import session_scope
    from app.models import User, Workbook
    with session_scope() as session:
        other = User(email="someone@else")
        session.add(other)
        session.flush()
        session.get(Workbook, wb["id"]).owner_id = other.id
    assert client.get(f"/api/workbooks/{wb['id']}").status_code == 404
    assert client.get(f"/api/workbooks/{wb['id']}/export",
                      params={"version_id": wb["active_version"]["id"]}).status_code == 404
    assert client.get(f"/api/jobs/{json.dumps('x')}").status_code == 404
