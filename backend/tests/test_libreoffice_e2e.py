"""Genuine end-to-end tests with the real LibreOffice calculation engine.

Skipped (and reported as skipped) when `soffice` is not installed. Compares
LibreOffice's recalculation of the workbook's own formulas with the trusted
outputs from the independent Python model.
"""

from __future__ import annotations

import io

import pytest
from openpyxl import load_workbook

from app.calc.worker import recalculate_bytes
from app.demo import generate
from app.demo.model import compute, scenario_inputs
from app.workbook.compare import Tolerance, values_equal
from tests.conftest import HAS_SOFFICE, H, load_demo, run_scenario

pytestmark = [pytest.mark.libreoffice,
              pytest.mark.skipif(not HAS_SOFFICE, reason="LibreOffice (soffice) is not installed")]
TOL = Tolerance()


def _demo_bytes() -> bytes:
    wb, cached = generate.build_workbook()
    buf = io.BytesIO()
    wb.save(buf)
    return generate.inject_cached_values(buf.getvalue(), cached)


def test_libreoffice_recalculates_edited_inputs():
    wb = load_workbook(io.BytesIO(_demo_bytes()))
    wb["Inputs"]["B4"] = 32.0 * 1.08
    buf = io.BytesIO()
    wb.save(buf)  # openpyxl drops cached values; LibreOffice must compute them
    out = recalculate_bytes(buf.getvalue(), soffice="soffice", timeout=120)
    values = load_workbook(io.BytesIO(out), data_only=True)
    expected = generate.reference_cells(compute(scenario_inputs(1.08)))
    mismatches = [k for k, v in expected.items()
                  if not values_equal(v, values[k.split("!")[0]][k.split("!")[1]].value, TOL)]
    assert mismatches == []


def test_libreoffice_forces_recalculation_over_stale_cached_values():
    """Proves the worker profile forces a hard recalc on load: a deliberately wrong stored result is replaced."""
    wb, cached = generate.build_workbook()
    buf = io.BytesIO()
    wb.save(buf)
    wrong = dict(cached)
    wrong[("Summary", "B6")] = 999.0
    wrong[("Products", "N5")] = -1.0
    out = recalculate_bytes(generate.inject_cached_values(buf.getvalue(), wrong), soffice="soffice", timeout=120)
    values = load_workbook(io.BytesIO(out), data_only=True)
    assert values["Summary"]["B6"].value == pytest.approx(cached[("Summary", "B6")], rel=1e-12)
    assert values["Products"]["N5"].value == pytest.approx(cached[("Products", "N5")], rel=1e-12)


def test_real_engine_full_workflow(make_client):
    client = make_client(fake=False)
    wb = load_demo(client)
    assert wb["compatibility_status"] == "editable", wb["compatibility"]
    scenario = run_scenario(client, wb, 1.08)
    assert scenario["status"] == "ready", scenario
    expected = {r["product"]: r["profit"] for r in compute(scenario_inputs(1.08))["products"]}
    for row in scenario["comparison"]["tables"][0]["rows"]:
        assert row["cells"]["profit"]["after"] == pytest.approx(expected[row["label"]], rel=1e-9)
    applied = client.post(f"/api/scenarios/{scenario['id']}/apply", headers={**H, "Idempotency-Key": "lo"},
                          json={"expected_active_version_id": wb["active_version"]["id"]}).json()["version"]
    exported = client.get(f"/api/workbooks/{wb['id']}/export", params={"version_id": applied["id"]})
    reopened = load_workbook(io.BytesIO(exported.content))
    assert reopened["Products"]["N5"].value == "=G5-M5"
    assert reopened["Products"]["O5"].value == "=IF(G5=0,0,N5/G5)"
    cached = load_workbook(io.BytesIO(exported.content), data_only=True)
    assert cached["Products"]["N5"].value == pytest.approx(expected["Walnut Side Table"], rel=1e-9)


def test_real_engine_reports_new_formula_errors(make_client):
    """Setting every product volume to zero makes overhead allocation divide by zero."""
    client = make_client(fake=False)
    wb = load_demo(client)
    new_map = {k: v for k, v in wb["map"].items() if k not in ("workbook_id", "mapping_version_id",
                                                                 "source_version_id")}
    for i in range(6):
        new_map["fields"].append({"id": f"volume_{i}", "label": f"Volume {i}", "sheet": "Products",
                                  "address": f"C{4 + i}", "role": "input", "value_type": "number",
                                  "unit": "units/year", "minimum": 0, "editable": True, "confirmed": True})
    r = client.put(f"/api/workbooks/{wb['id']}/map", headers=H,
                   json={"expected_mapping_version_id": wb["mapping_version_id"], "map": new_map})
    assert r.status_code == 200, r.text
    r = client.post(f"/api/workbooks/{wb['id']}/scenarios", headers=H, json={
        "base_version_id": wb["active_version"]["id"], "name": "zero volume",
        "changes": [{"field_id": f"volume_{i}", "operation": "set", "operand": 0} for i in range(6)]})
    from tests.conftest import wait_scenario
    scenario = wait_scenario(client, r.json()["scenario"]["id"], timeout=180)
    assert scenario["status"] == "failed"
    assert any(v["code"] == "no_new_formula_errors" and not v["passed"] for v in scenario["validations"])
