"""Shared fixtures.

Most workflow tests use `fake_recalculate`, a test double for the LibreOffice
worker that only understands the synthetic demo workbook: it reads the inputs,
evaluates them with the independent Python model, and stores the results as
cached values. Tests marked `libreoffice` use the real engine instead.
"""

from __future__ import annotations

import io
import shutil
import time
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.config import Settings
from app.demo import generate
from app.demo.model import Inputs, Product, compute

H = {"X-Requested-With": "projections-assistant"}
HAS_SOFFICE = shutil.which("soffice") is not None


def fake_recalculate(data: bytes, *, soffice: str, timeout: float, work_root=None) -> bytes:
    wb = load_workbook(io.BytesIO(data))
    inp = wb["Inputs"]
    inputs = Inputs(manufacturing_hourly_wage=inp["B4"].value, packaging_hourly_wage=inp["B5"].value,
                    payroll_burden_rate=inp["B6"].value, annual_fixed_overhead=inp["B7"].value)
    ws = wb["Products"]
    products = tuple(Product(ws[f"A{r}"].value, ws[f"B{r}"].value, ws[f"C{r}"].value, ws[f"D{r}"].value,
                             ws[f"E{r}"].value, ws[f"F{r}"].value)
                     for r in range(generate.FIRST_ROW, generate.LAST_ROW + 1))
    cells = generate.reference_cells(compute(inputs, products))
    cached = {tuple(k.split("!", 1)): v for k, v in cells.items()}
    out = io.BytesIO()
    wb.save(out)
    return generate.inject_cached_values(out.getvalue(), cached)


def make_settings(tmp_path: Path, **overrides) -> Settings:
    values = {"data_dir": tmp_path / "data", "worker_mode": "inprocess", "calc_wait_in_turn_seconds": 15.0,
              "_env_file": None}
    values.update(overrides)
    return Settings(**values)


@pytest.fixture
def fake_worker(monkeypatch):
    import app.calc.worker as worker
    monkeypatch.setattr(worker, "recalculate_bytes", fake_recalculate)


@pytest.fixture
def make_client(tmp_path, monkeypatch):
    clients = []

    def _make(*, fake: bool = True, anthropic_client=None, access_verifier=None, **overrides):
        if fake:
            import app.calc.worker as worker
            monkeypatch.setattr(worker, "recalculate_bytes", fake_recalculate)
        from app.main import create_app
        app = create_app(make_settings(tmp_path, **overrides), anthropic_client=anthropic_client,
                         access_verifier=access_verifier)
        client = TestClient(app)
        client.__enter__()
        clients.append(client)
        return client

    yield _make
    for c in clients:
        c.__exit__(None, None, None)


def wait_job(client: TestClient, job_id: str, timeout: float = 60) -> dict:
    deadline = time.time() + timeout
    while time.time() < deadline:
        job = client.get(f"/api/jobs/{job_id}").json()
        if job["status"] in ("succeeded", "failed"):
            return job
        time.sleep(0.2)
    raise AssertionError(f"job {job_id} did not finish")


def wait_scenario(client: TestClient, scenario_id: str, timeout: float = 60) -> dict:
    deadline = time.time() + timeout
    while time.time() < deadline:
        s = client.get(f"/api/scenarios/{scenario_id}").json()
        if s["status"] not in ("draft", "validating", "calculating"):
            return s
        time.sleep(0.2)
    raise AssertionError(f"scenario {scenario_id} did not finish")


def load_demo(client: TestClient) -> dict:
    r = client.post("/api/workbooks/demo", headers=H)
    assert r.status_code == 201, r.text
    body = r.json()
    job = wait_job(client, body["job_id"])
    assert job["status"] == "succeeded", job
    wb = client.get(f"/api/workbooks/{body['workbook']['id']}").json()
    return wb


def run_scenario(client: TestClient, wb: dict, factor: float, *, base: str | None = None,
                 revision_of: str | None = None) -> dict:
    r = client.post(f"/api/workbooks/{wb['id']}/scenarios", headers=H, json={
        "base_version_id": base or wb["active_version"]["id"], "name": f"wage x{factor}",
        "changes": [{"field_id": "manufacturing_hourly_rate", "operation": "multiply", "operand": factor}],
        "revision_of_scenario_id": revision_of})
    assert r.status_code == 201, r.text
    return wait_scenario(client, r.json()["scenario"]["id"])
